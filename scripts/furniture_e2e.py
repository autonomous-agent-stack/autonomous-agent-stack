from __future__ import annotations

import hashlib
import json
import struct
import zlib
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from autoresearch.core.services.approval_store import ApprovalStoreService
from autoresearch.core.services.governed_mcp import (
    GovernedMCPService,
    GovernedMCPToolCallRequest,
    ToolPermissionService,
)
from autoresearch.core.services.model_gateway import ModelGatewayService, ModelProviderRead
from autoresearch.core.services.session_events import SessionEventService
from autoresearch.core.services.usage_quota import UsageLedgerEntryRead, UsageQuotaService
from autoresearch.ga.contracts import (
    ImageGenerationRequest,
    ModelPolicyRead,
    PolicyDecisionRead,
    PolicyDecisionValue,
    PrincipalRead,
    PrincipalType,
)
from autoresearch.shared.models import (
    ApprovalDecisionRequest,
    ApprovalRequestCreateRequest,
    ApprovalRequestRead,
    ApprovalStatus,
    SessionEventCreateRequest,
)
from autoresearch.shared.store import InMemoryRepository
from autoresearch.storage.events import InMemorySessionEventStore, verify_event_hash


ARTIFACT_ROOT = Path("artifacts/ga/furniture_e2e")
REQUIRED_ARTIFACTS = {
    "quote.pdf",
    "quote.xlsx",
    "design_image.png",
    "design_prompt_audit.json",
    "approval_record.json",
    "audit_timeline.json",
    "cost_ledger.json",
    "commission_projection.json",
    "session_facts_replay.json",
}


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    runner = FurnitureE2ERunner(repo_root)
    manifest = runner.run()
    print(json.dumps({"status": "passed", "artifact_root": str(ARTIFACT_ROOT), "artifacts": manifest}, sort_keys=True))
    return 0


class FurnitureE2ERunner:
    def __init__(self, repo_root: Path) -> None:
        self.repo_root = repo_root
        self.artifact_root = repo_root / ARTIFACT_ROOT
        self.artifact_root.mkdir(parents=True, exist_ok=True)
        self.session_id = "furniture-e2e-ga-session"
        self.session_events = SessionEventService(event_store=InMemorySessionEventStore())
        self.approvals = ApprovalStoreService(
            repository=InMemoryRepository[ApprovalRequestRead](),
            session_events=self.session_events,
        )
        self.gateway = ModelGatewayService(
            providers=[ModelProviderRead(provider_id="local-dev", enabled=True, models=["noop", "image2"])]
        )
        self.mcp = GovernedMCPService(
            servers_path=repo_root / "configs/mcp_servers.yaml",
            permission_service=ToolPermissionService(policy_path=repo_root / "configs/tool_permissions.yaml"),
            quota_service=UsageQuotaService(
                repository=InMemoryRepository[UsageLedgerEntryRead](),
                policy_path=repo_root / "configs/quota_policy.yaml",
            ),
            approval_store=self.approvals,
            session_events=self.session_events,
        )

    def run(self) -> dict[str, Any]:
        self._append("customer.inquiry.received", "Customer requested a modular walnut workstation quote.")
        quote = _quote_payload()
        self._write_quote_pdf(quote)
        self._write_quote_xlsx(quote)
        image_result = self._generate_design_image()
        approval_record = self._exercise_approval_paths()
        cost_ledger = self._write_cost_ledger(quote, image_result)
        self._write_commission_projection(quote)
        self._append(
            "quote.artifacts.generated",
            "Quote PDF, workbook, governed design image, cost ledger, and commission projection generated.",
            artifact_refs=[
                {"artifact_type": "quote_pdf", "path": str(ARTIFACT_ROOT / "quote.pdf")},
                {"artifact_type": "quote_xlsx", "path": str(ARTIFACT_ROOT / "quote.xlsx")},
                {"artifact_type": "design_image", "path": str(ARTIFACT_ROOT / "design_image.png")},
            ],
            payload={"quote_id": quote["quote_id"], "cost_total": cost_ledger["totals"]["total_cost"]},
        )
        self._write_json("approval_record.json", approval_record)
        self._write_audit_timeline()
        self._write_session_facts_replay()
        self._validate_outputs()
        return {name: str(ARTIFACT_ROOT / name) for name in sorted(REQUIRED_ARTIFACTS)}

    def _generate_design_image(self) -> dict[str, Any]:
        principal = PrincipalRead(
            principal_id="furniture-design-agent",
            principal_type=PrincipalType.SERVICE_ACCOUNT,
            roles=["sales", "designer"],
        )
        decision = PolicyDecisionRead(
            decision_id="pd-furniture-image-allow",
            decision=PolicyDecisionValue.ALLOW,
            subject=principal,
            action="model.image.generate",
            resource="model://local-dev/image2",
            reason="approved furniture E2E dry-run image generation",
        )
        result = self.gateway.generate_image(
            ImageGenerationRequest(
                provider_id="local-dev",
                model_id="image2",
                prompt=(
                    "Render a walnut modular workstation for customer Mei at mei@example.com, "
                    "phone +1 415 555 0132, with cable management and warm neutral lighting."
                ),
                principal=principal,
                session_id=self.session_id,
                policy_decision_id=decision.decision_id,
                policy=ModelPolicyRead(
                    policy_id="model-policy-furniture-image",
                    allowed_modalities=["image"],
                    cost_center="furniture-sales",
                ),
                metadata={"dry_run": True, "workflow": "furniture-e2e"},
            ),
            policy_decision=decision,
        )
        self.session_events.append(SessionEventCreateRequest.model_validate(result.session_event))
        image_path = self.artifact_root / "design_image.png"
        image_path.write_bytes(_render_png(result.artifact.content_hash))
        design_prompt_audit = {
            "generated_at": _now(),
            "model_policy": {"policy_id": result.design_prompt_audit.policy_id, "modality": "image"},
            "model_usage_ledger": result.usage_ledger.model_dump(mode="json"),
            "image_generation_artifact": {
                **result.artifact.model_dump(mode="json"),
                "path": str(ARTIFACT_ROOT / "design_image.png"),
                "file_sha256": _file_sha256(image_path),
            },
            "design_prompt_audit": result.design_prompt_audit.model_dump(mode="json"),
            "artifact_ref": result.session_event.get("artifact_refs", [])[0],
            "session_event": result.session_event,
            "pii_redaction_verified": result.design_prompt_audit.pii_redacted
            and "mei@example.com" not in result.design_prompt_audit.redacted_prompt,
        }
        self._write_json("design_prompt_audit.json", design_prompt_audit)
        return design_prompt_audit

    def _exercise_approval_paths(self) -> dict[str, Any]:
        denied_first = self.mcp.call_tool(
            GovernedMCPToolCallRequest(
                tool_id="local_demo.external_write",
                params={"recipient": "customer@example.com", "message": "quote draft"},
                actor_id="sales-agent",
                actor_role="supervisor",
                session_id=self.session_id,
                metadata={"workflow": "furniture-e2e-denied"},
            )
        )
        if not denied_first.approval_id:
            raise RuntimeError("denied path did not create approval")
        denied_approval = self.approvals.resolve_request(
            denied_first.approval_id,
            ApprovalDecisionRequest(decision="rejected", decided_by="ga-furniture-e2e", note="deny path drill"),
        )
        denied_followup = self.mcp.call_tool(
            GovernedMCPToolCallRequest(
                tool_id="local_demo.external_write",
                params={"recipient": "customer@example.com", "message": "quote draft"},
                actor_id="sales-agent",
                actor_role="supervisor",
                session_id=self.session_id,
                approval_id=denied_approval.approval_id,
                metadata={"workflow": "furniture-e2e-denied-followup"},
            )
        )

        approved_first = self.mcp.call_tool(
            GovernedMCPToolCallRequest(
                tool_id="local_demo.external_write",
                params={"recipient": "customer@example.com", "message": "approved quote dry-run"},
                actor_id="sales-agent",
                actor_role="supervisor",
                session_id=self.session_id,
                metadata={"workflow": "furniture-e2e-approved"},
            )
        )
        if not approved_first.approval_id:
            raise RuntimeError("approved path did not create approval")
        approved_approval = self.approvals.resolve_request(
            approved_first.approval_id,
            ApprovalDecisionRequest(decision="approved", decided_by="ga-furniture-e2e", note="approved dry-run path"),
        )
        approved_write = self.mcp.call_tool(
            GovernedMCPToolCallRequest(
                tool_id="local_demo.external_write",
                params={"recipient": "customer@example.com", "message": "approved quote dry-run"},
                actor_id="sales-agent",
                actor_role="supervisor",
                session_id=self.session_id,
                approval_id=approved_approval.approval_id,
                metadata={
                    "workflow": "furniture-e2e-approved-write",
                    "policy_decision": "allow",
                    "policy_decision_id": "pd-furniture-external-write",
                    "recipient_allowlist": ["customer@example.com"],
                    "audit_timeline_id": "audit-furniture-e2e",
                    "session_facts_id": "facts-furniture-e2e",
                    "live_credentials": False,
                },
            )
        )
        return {
            "generated_at": _now(),
            "denied_path": {
                "approval": denied_approval.model_dump(mode="json"),
                "initial_status": denied_first.status,
                "followup_status": denied_followup.status,
                "blocked": denied_followup.status != "succeeded",
            },
            "approved_path": {
                "approval": approved_approval.model_dump(mode="json"),
                "initial_status": approved_first.status,
                "write_status": approved_write.status,
                "dry_run": approved_write.result.get("dry_run") is True,
                "result": approved_write.model_dump(mode="json"),
            },
        }

    def _write_quote_pdf(self, quote: dict[str, Any]) -> None:
        lines = [
            "AAS Furniture Quote",
            f"Quote: {quote['quote_id']}",
            f"Customer: {quote['customer']['company']}",
            f"Items: {len(quote['items'])}",
            f"Subtotal: ${quote['totals']['subtotal']:.2f}",
            f"Total: ${quote['totals']['total']:.2f}",
            "Status: governed dry-run external write",
        ]
        (self.artifact_root / "quote.pdf").write_bytes(_minimal_pdf(lines))

    def _write_quote_xlsx(self, quote: dict[str, Any]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Quote"
        sheet.append(["Quote ID", quote["quote_id"]])
        sheet.append(["Customer", quote["customer"]["company"]])
        sheet.append([])
        sheet.append(["SKU", "Description", "Qty", "Unit Price", "Line Total"])
        for item in quote["items"]:
            sheet.append([item["sku"], item["description"], item["qty"], item["unit_price"], item["line_total"]])
        sheet.append([])
        sheet.append(["Subtotal", quote["totals"]["subtotal"]])
        sheet.append(["Tax", quote["totals"]["tax"]])
        sheet.append(["Total", quote["totals"]["total"]])
        workbook.save(self.artifact_root / "quote.xlsx")

    def _write_cost_ledger(self, quote: dict[str, Any], image_result: dict[str, Any]) -> dict[str, Any]:
        ledger = {
            "generated_at": _now(),
            "quote_id": quote["quote_id"],
            "currency": "USD",
            "line_items": [
                {"kind": "materials", "amount": 1840.0},
                {"kind": "labor", "amount": 420.0},
                {"kind": "logistics", "amount": 260.0},
                {
                    "kind": "model_gateway_image",
                    "amount": image_result["model_usage_ledger"]["cost_units"],
                    "usage_id": image_result["model_usage_ledger"]["usage_id"],
                },
            ],
        }
        ledger["totals"] = {"total_cost": sum(float(item["amount"]) for item in ledger["line_items"])}
        self._write_json("cost_ledger.json", ledger)
        return ledger

    def _write_commission_projection(self, quote: dict[str, Any]) -> None:
        projection = {
            "generated_at": _now(),
            "quote_id": quote["quote_id"],
            "sales_owner": "furniture-sales-agent",
            "basis": quote["totals"]["subtotal"],
            "rate": 0.05,
            "projected_commission": round(quote["totals"]["subtotal"] * 0.05, 2),
            "approval_required_before_external_write": True,
        }
        self._write_json("commission_projection.json", projection)

    def _write_audit_timeline(self) -> None:
        timeline = self.session_events.timeline(session_id=self.session_id, limit=1000)
        self._write_json("audit_timeline.json", timeline.model_dump(mode="json"))

    def _write_session_facts_replay(self) -> None:
        events = self.session_events.list_events(session_id=self.session_id, limit=1000)
        replay = {
            "generated_at": _now(),
            "session_id": self.session_id,
            "event_count": len(events),
            "content_hash_verified": all(verify_event_hash(event) for event in events),
            "events": [event.model_dump(mode="json") for event in events],
        }
        self._write_json("session_facts_replay.json", replay)

    def _append(
        self,
        event_type: str,
        content: str,
        *,
        payload: dict[str, Any] | None = None,
        artifact_refs: list[dict[str, Any]] | None = None,
    ) -> None:
        self.session_events.append(
            SessionEventCreateRequest(
                session_id=self.session_id,
                source="furniture_e2e",
                event_type=event_type,
                content=content,
                payload=dict(payload or {}),
                artifact_refs=list(artifact_refs or []),
                idempotency_key=f"furniture-e2e:{event_type}",
            )
        )

    def _write_json(self, name: str, payload: dict[str, Any]) -> None:
        (self.artifact_root / name).write_text(
            json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )

    def _validate_outputs(self) -> None:
        missing = [name for name in sorted(REQUIRED_ARTIFACTS) if not (self.artifact_root / name).exists()]
        if missing:
            raise RuntimeError(f"missing furniture artifacts: {', '.join(missing)}")
        empty = [name for name in sorted(REQUIRED_ARTIFACTS) if (self.artifact_root / name).stat().st_size <= 0]
        if empty:
            raise RuntimeError(f"empty furniture artifacts: {', '.join(empty)}")
        if not (self.artifact_root / "quote.pdf").read_bytes().startswith(b"%PDF-"):
            raise RuntimeError("quote.pdf is not a PDF")
        load_workbook(self.artifact_root / "quote.xlsx", read_only=True).close()
        if not (self.artifact_root / "design_image.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n"):
            raise RuntimeError("design_image.png is not a PNG")
        approval = _read_json(self.artifact_root / "approval_record.json")
        if approval["denied_path"]["blocked"] is not True:
            raise RuntimeError("approval-denied path did not block continuation")
        if approval["approved_path"]["write_status"] != "succeeded" or approval["approved_path"]["dry_run"] is not True:
            raise RuntimeError("approval-approved path did not produce dry-run external write")
        design = _read_json(self.artifact_root / "design_prompt_audit.json")
        if design["pii_redaction_verified"] is not True:
            raise RuntimeError("design prompt PII redaction was not verified")
        replay = _read_json(self.artifact_root / "session_facts_replay.json")
        if replay["content_hash_verified"] is not True or replay["event_count"] < 6:
            raise RuntimeError("session facts replay did not verify event hashes")


def _quote_payload() -> dict[str, Any]:
    items = [
        {"sku": "WAL-WS-180", "description": "Walnut workstation 180cm", "qty": 4, "unit_price": 1280.0},
        {"sku": "CAB-MOD-02", "description": "Modular storage cabinet", "qty": 4, "unit_price": 420.0},
        {"sku": "LED-WARM-01", "description": "Warm task lighting kit", "qty": 4, "unit_price": 180.0},
    ]
    for item in items:
        item["line_total"] = round(float(item["qty"]) * float(item["unit_price"]), 2)
    subtotal = round(sum(float(item["line_total"]) for item in items), 2)
    tax = round(subtotal * 0.0825, 2)
    return {
        "quote_id": "FURN-GA-0001",
        "customer": {"company": "Northstar Studio", "contact_ref": "customer-ref-001"},
        "items": items,
        "totals": {"subtotal": subtotal, "tax": tax, "total": round(subtotal + tax, 2)},
    }


def _minimal_pdf(lines: list[str]) -> bytes:
    escaped = [_pdf_text(line) for line in lines]
    content_lines = ["BT", "/F1 16 Tf", "72 742 Td"]
    for index, line in enumerate(escaped):
        if index:
            content_lines.append("0 -24 Td")
        content_lines.append(f"({line}) Tj")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = b"%PDF-1.4\n"
    offsets = [0]
    for obj_no, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{obj_no} 0 obj\n".encode("ascii") + obj + b"\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    pdf += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010d} 00000 n \n".encode("ascii")
    pdf += (
        b"trailer\n"
        + f"<< /Root 1 0 R /Size {len(objects) + 1} >>\n".encode("ascii")
        + b"startxref\n"
        + str(xref_offset).encode("ascii")
        + b"\n%%EOF\n"
    )
    return pdf


def _render_png(seed: str, *, width: int = 160, height: int = 96) -> bytes:
    digest = hashlib.sha256(seed.encode("utf-8")).digest()
    rows = []
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            band = (x * 3 + y * 5) % 256
            row.extend(
                [
                    (digest[0] + band) % 256,
                    (digest[8] + x * 2) % 256,
                    (digest[16] + y * 3) % 256,
                ]
            )
        rows.append(bytes(row))
    raw = b"".join(rows)
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + _png_chunk(b"IDAT", zlib.compress(raw))
        + _png_chunk(b"IEND", b"")
    )


def _png_chunk(kind: bytes, data: bytes) -> bytes:
    crc = zlib.crc32(kind + data) & 0xFFFFFFFF
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", crc)


def _pdf_text(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_json(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise RuntimeError(f"{path.name} is not a JSON object")
    return payload


def _now() -> str:
    return datetime.now(UTC).isoformat()


if __name__ == "__main__":
    raise SystemExit(main())
